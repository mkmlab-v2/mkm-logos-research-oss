#!/usr/bin/env python3
"""One-off probe: section table rows in filled vs template docx."""
from __future__ import annotations

from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
FILLED = ROOT / "reports/kstartup_startup_package_ai_filled/doyak_plan_filled_v1.docx"
TEMPLATE = ROOT / "reports/kstartup_startup_package_ai_attachments/(도약) [별첨 1] 2026년 창업도약패키지(AI 인재 실증형) 사업계획서 양식.docx"
BLUE = {"0000FF", "0070C0", "4472C4", "2F5496", "0563C1", "5B9BD5"}


def blue_in_cell(cell) -> list[str]:
    out = []
    for p in cell.paragraphs:
        for run in p.runs:
            col = run.font.color
            if col and col.rgb and str(col.rgb).upper() in BLUE and (run.text or "").strip():
                out.append(run.text.strip()[:100])
    return out


def probe(path: Path, label: str) -> None:
    doc = Document(str(path))
    print(f"\n=== {label} {path.name} ===")
    for ti, table in enumerate(doc.tables):
        blob = "".join(c.text for r in table.rows for c in r.cells)
        if "아이템 개요" not in blob and "문제 인식" not in blob:
            continue
        print(f"section_table T{ti} rows={len(table.rows)} cols={len(table.rows[0].cells) if table.rows else 0}")
        for ri, row in enumerate(table.rows):
            c0 = (row.cells[0].text or "").replace("\n", " | ")
            c1 = (row.cells[1].text if len(row.cells) > 1 else "").replace("\n", " ")[:120]
            bl0 = blue_in_cell(row.cells[0])
            bl1 = blue_in_cell(row.cells[1]) if len(row.cells) > 1 else []
            if not c0.strip() and not c1.strip() and not bl0 and not bl1:
                continue
            print(f"  r{ri:02d} c0={c0[:70]!r}")
            print(f"       c1_len={len(row.cells[1].text) if len(row.cells)>1 else 0} c1={c1!r}")
            if bl0:
                print(f"       BLUE_c0={bl0[:3]}")
            if bl1:
                print(f"       BLUE_c1={bl1[:3]}")


def scan_instructions(path: Path) -> None:
    doc = Document(str(path))
    keys = (
        "※",
        "창업 아이템의 필요성",
        "개발 계획",
        "개발 아이템",
        "< 사진",
        "세부내용 작성",
    )
    print(f"\n--- instruction scan {path.name} ---")
    for ti, table in enumerate(doc.tables):
        seen: set[int] = set()
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                tc = id(cell._tc)
                if tc in seen:
                    continue
                seen.add(tc)
                tx = cell.text or ""
                if any(k in tx for k in keys):
                    bl = blue_in_cell(cell)
                    print(
                        f"T{ti}r{ri}c{ci} len={len(tx)} blue={len(bl)} "
                        f"preview={tx[:120].replace(chr(10), ' | ')!r}"
                    )
    for i, p in enumerate(doc.paragraphs):
        t = p.text or ""
        if any(k in t for k in keys):
            bl = [
                r.text[:60]
                for r in p.runs
                if r.font.color
                and r.font.color.rgb
                and str(r.font.color.rgb).upper() in BLUE
                and (r.text or "").strip()
            ]
            print(f"P{i} blue_runs={len(bl)} {t[:140]!r}")


if __name__ == "__main__":
    if TEMPLATE.is_file():
        probe(TEMPLATE, "TEMPLATE")
    probe(FILLED, "FILLED")
    scan_instructions(FILLED)
