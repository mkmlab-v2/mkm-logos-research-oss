#!/usr/bin/env python3
"""Fill 도약 별첨1 docx from paste_ready pack and export PDF (Word COM).

Outputs:
  reports/kstartup_startup_package_ai_filled/doyak_plan_filled_v1.docx
  reports/kstartup_startup_package_ai_filled/doyak_plan_filled_v1.pdf
  reports/kstartup_startup_package_ai_filled/ai_talent_2p_filled_v1.pdf
  reports/kstartup_startup_package_ai_filled_plan_latest.json
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]

from kstartup_startup_package_ai_form_notice_compliance_v1 import (  # noqa: E402
    TEAM_CAPABILITY_DETAIL,
    TEAM_CAPABILITY_GENERAL,
    ai_talent_body_summary,
    apply_form_notice_compliance,
)
CONFIG = ROOT / "docs/final/artifacts/kstartup_startup_package_ai_pms_config_v1.json"
ELIGIBILITY = ROOT / "docs/final/artifacts/startup_package_ai_2026_eligibility_v1_latest.json"
PASTE_DIR = ROOT / "reports/kstartup_startup_package_ai_paste_ready"
OUT_DIR = ROOT / "reports/kstartup_startup_package_ai_filled"
OUT_DOCX = OUT_DIR / "doyak_plan_filled_v1.docx"
OUT_PDF = OUT_DIR / "doyak_plan_filled_v1.pdf"
OUT_AI_PDF = OUT_DIR / "ai_talent_2p_filled_v1.pdf"
OUT_META = ROOT / "reports/kstartup_startup_package_ai_filled_plan_latest.json"

SECTION_MAP: list[tuple[str, str]] = [
    ("아이템 개요", "plan_01_summary_paste.txt"),
    ("문제 인식", "plan_02_market_problem_paste.txt"),
    ("실현 가능성", "plan_03_tech_roadmap_paste.txt"),
    ("성장전략", "plan_04_growth_funding_paste.txt"),
    ("팀 구성", "plan_05_team_paste.txt"),
    ("AI 인재 활용", "plan_06_ai_talent_2p_paste.txt"),
]

TOC_MARKERS = ("사업계획서 작성 목차", "항목", "세부항목")


def _utc() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _load_config() -> dict[str, Any]:
    if CONFIG.is_file():
        return json.loads(CONFIG.read_text(encoding="utf-8-sig"))
    return {}


def _load_eligibility() -> dict[str, Any]:
    if ELIGIBILITY.is_file():
        return json.loads(ELIGIBILITY.read_text(encoding="utf-8-sig"))
    return {}


def _entity_profile(cfg: dict[str, Any], elig: dict[str, Any]) -> dict[str, Any]:
    ent = dict(cfg.get("legal_entity") or {})
    leg = elig.get("legal_entity") or {}
    for k, v in leg.items():
        if k == "incorporation_facts":
            continue
        if v and not ent.get(k):
            ent[k] = v
    facts = leg.get("incorporation_facts") or {}
    if facts.get("개업연월일"):
        ent["founding_date"] = str(facts["개업연월일"]).replace("-", ".")
    return ent


def _paste_body(path: Path) -> str:
    raw = path.read_text(encoding="utf-8")
    raw = re.sub(r"^\[창업패키지[^\]]*\]\s*\n+", "", raw)
    raw = re.sub(r"^#+ .+\n+", "", raw, count=1)
    return _clean_paste_for_word(raw.strip())


def _clean_paste_for_word(text: str) -> str:
    """Strip markdown tables/separators so Word cells stay readable."""
    lines: list[str] = []
    for line in text.splitlines():
        s = line.strip()
        if re.match(r"^\|[-: |]+\|$", s):
            continue
        if s.startswith("|") and s.endswith("|"):
            s = re.sub(r"^\||\|$", "", s)
            s = " · ".join(p.strip() for p in s.split("|") if p.strip())
        if s in ("---", "—", "***"):
            continue
        lines.append(s)
    out = "\n".join(lines)
    out = re.sub(r"\n{3,}", "\n\n", out)
    return out.strip()


def _load_pastes() -> dict[str, str]:
    out: dict[str, str] = {}
    for _label, fname in SECTION_MAP:
        p = PASTE_DIR / fname
        if p.is_file():
            body = _paste_body(p)
            out[fname] = apply_form_notice_compliance(body, context=fname)
    return out


BLUE_RGB = frozenset(
    {"0000FF", "0070C0", "4472C4", "2F5496", "0563C1", "5B9BD5", "2E75B6"}
)


def _run_is_blue(run) -> bool:
    col = run.font.color
    if col is not None and col.rgb is not None:
        return str(col.rgb).upper() in BLUE_RGB
    return False


def _apply_black_font(run) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.name = "맑은 고딕"
    run.font.size = Pt(10)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "맑은 고딕")
    rfonts.set(qn("w:hAnsi"), "맑은 고딕")
    rfonts.set(qn("w:eastAsia"), "맑은 고딕")


def _clear_cell_content(cell) -> None:
    from docx.oxml.ns import qn

    tc = cell._tc
    for tbl in list(tc.findall(qn("w:tbl"))):
        tc.remove(tbl)
    for p_el in list(tc.findall(qn("w:p"))):
        tc.remove(p_el)


def _set_cell_text(cell, text: str, *, max_len: int = 14_000) -> int:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    body = (text or "").strip()[:max_len]
    tc = cell._tc
    for p_el in list(tc.findall(qn("w:p"))):
        tc.remove(p_el)
    if not body:
        cell.add_paragraph()
        return 0
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(body)
    _apply_black_font(run)
    return len(body)


def _set_section_cell_content(cell, text: str, *, max_len: int = 14_000) -> int:
    """Section 본문 칸: 중첩 안내 표(파란 ※) 제거 후 단락별 검정 본문."""
    from docx.shared import Pt

    body = (text or "").strip()[:max_len]
    _clear_cell_content(cell)
    if not body:
        cell.add_paragraph()
        return 0
    total = 0
    for block in re.split(r"\n\n+", body):
        block = block.strip()
        if not block:
            continue
        for line in block.split("\n"):
            line = line.strip()
            if not line:
                continue
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(line)
            _apply_black_font(run)
            total += len(line)
    return total


def _set_span(row, col_start: int, col_end: int, text: str) -> None:
    """Write once into the first logical cell (merged-safe)."""
    if col_start < len(row.cells):
        _set_cell_text(row.cells[col_start], text)


def _check_cell(cell) -> None:
    tx = (cell.text or "").strip()
    if tx == "□":
        _set_cell_text(cell, "■")


INSTRUCTION_RUN_RE = re.compile(
    r"(법인등기부등본|사업자등록증|본사\(점\)|동일하게\s*기입|기준으로\s*기입|"
    r"OOOOOO-OOOOOOO|OO도\s*OO시|^\(※|※\s*사업\s*신청|※\s*정부지원|"
    r"개발하고자|세부내용\s*작성|유추가능한\s*정보|평가항목입니다)"
)

INSTRUCTION_PARA_RE = re.compile(
    r"^(※|ㅇ\s*AI\s|ㅇ\s*전체|<\s*사업|<\s*팀|<\s*협력|<\s*예시|<\s*AI)"
)

SECTION_HEADER_RE = re.compile(
    r"(1\.\s*문제\s*인식|2\.\s*실현\s*가능성|3\.\s*성장전략|창업\s*아이템\s*개요)"
)


def _paragraph_xml_text(paragraph) -> str:
    return "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", paragraph._element.xml))


def _paragraph_has_blue_xml(paragraph) -> bool:
    xml = paragraph._element.xml.upper()
    return any(rgb in xml for rgb in BLUE_RGB)


def _is_instruction_body_paragraph(paragraph) -> bool:
    xt = _paragraph_xml_text(paragraph).strip()
    if not xt:
        return False
    if xt.startswith("※") or "※" in xt:
        if _paragraph_has_blue_xml(paragraph) or INSTRUCTION_RUN_RE.search(xt):
            return True
    t = (paragraph.text or "").strip()
    return bool(t.startswith("※"))


def _clear_paragraph_runs(paragraph) -> None:
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)


def _set_paragraph_black_body(paragraph, text: str, *, max_len: int = 6_000) -> int:
    from docx.shared import Pt

    body = (text or "").strip()[:max_len]
    _clear_paragraph_runs(paragraph)
    if not body:
        return 0
    run = paragraph.add_run(body)
    _apply_black_font(run)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.space_before = Pt(0)
    return len(body)


def _numbered_paste_blocks(text: str) -> dict[str, str]:
    blocks: dict[str, list[str]] = {}
    current: str | None = None
    for line in text.splitlines():
        m = re.match(r"^(\d+\.\d+)\s+", line.strip())
        if m:
            current = m.group(1)
            blocks.setdefault(current, []).append(line)
        elif current:
            blocks[current].append(line)
    return {k: "\n".join(v).strip() for k, v in blocks.items()}


def _problem_body_answers(pastes: dict[str, str]) -> list[str]:
    p02 = pastes.get("plan_02_market_problem_paste.txt", "")
    p01 = pastes.get("plan_01_summary_paste.txt", "")
    blocks = _numbered_paste_blocks(p02)
    need = "\n\n".join(x for x in (blocks.get("2.1", ""), blocks.get("2.2", "")) if x).strip()
    dev = blocks.get("2.3", "").strip()
    if not dev:
        dev = (
            "근거·서식·버전이 결합된 라스트마일 문서 파이프라인 없이는 "
            "공공 신청 품질·감사 추적을 안정적으로 확보하기 어렵다."
        )
    intro = _clean_paste_for_word(p01)
    intro = re.sub(r"^-\s*과제명\(안\):.*\n?", "", intro, flags=re.M).strip()
    return [
        f"○ 창업 아이템의 필요성 (국내·외 시장 현황 및 문제점)\n{need or p02[:900]}",
        f"○ 문제 해결을 위한 개발 필요성\n{dev}",
        f"○ 개발 아이템 소개\n{intro or p01[:700]}",
    ]


def _solution_body_answer(pastes: dict[str, str]) -> str:
    p03 = pastes.get("plan_03_tech_roadmap_paste.txt", "")
    blocks = _numbered_paste_blocks(p03)
    parts = [blocks.get(k, "") for k in ("3.1", "3.2", "3.3") if blocks.get(k)]
    body = "\n\n".join(parts).strip() or p03[:2_400]
    return f"○ 창업 아이템의 개발 계획 및 차별성\n{body}"


def _scale_body_answers(pastes: dict[str, str]) -> list[str]:
    p04 = pastes.get("plan_04_growth_funding_paste.txt", "")
    blocks = _numbered_paste_blocks(p04)
    m1 = "\n\n".join(x for x in (blocks.get("4.1", ""), blocks.get("4.2", "")) if x).strip()
    m2 = blocks.get("4.3", "").strip() or p04[-600:].strip()
    return [
        f"○ 목표 시장 진입·사업화 전략\n{m1 or p04[:900]}",
        f"○ 비즈니스·자금 확보 전략\n{m2}",
    ]


def _fill_body_flow_answers(doc, pastes: dict[str, str]) -> dict[str, Any]:
    """양식 본문 단락(표 밖): 파란 ※ 아래 ㅇ 칸에 검정 답변."""
    queues = {
        "problem": _problem_body_answers(pastes),
        "solution": [_solution_body_answer(pastes)],
        "scale": _scale_body_answers(pastes),
    }
    state: str | None = None
    idx = 0
    filled = {"problem": 0, "solution": 0, "scale": 0, "overview": 0}
    overview_pending = False

    for p in list(doc.paragraphs):
        xt = _paragraph_xml_text(p)
        t = (p.text or "").strip()

        if "창업 아이템 개요" in xt and "□" in xt:
            state = None
            overview_pending = True
            if t.startswith("□"):
                _set_paragraph_black_body(p, t.replace("□", "■", 1))
            continue

        if overview_pending and not t and not xt.strip():
            n = _set_paragraph_black_body(
                p, pastes.get("plan_01_summary_paste.txt", "")[:2_000]
            )
            if n:
                filled["overview"] = 1
            overview_pending = False
            continue

        if SECTION_HEADER_RE.search(xt.replace("_", " ")):
            if "문제 인식" in xt:
                state = "problem"
            elif "실현 가능성" in xt:
                state = "solution"
            elif "성장전략" in xt:
                state = "scale"
            else:
                state = None
            idx = 0
            overview_pending = False
            for run in p.runs:
                _apply_black_font(run)
            continue

        if state and re.fullmatch(r"ㅇ\s*", t):
            q = queues.get(state, [])
            if idx < len(q):
                if _set_paragraph_black_body(p, q[idx]):
                    filled[state] = int(filled.get(state, 0)) + 1
                idx += 1
            continue

    return filled


def _purge_blue_instruction_body_paragraphs(doc) -> int:
    removed = 0
    for p in list(doc.paragraphs):
        if _is_instruction_body_paragraph(p):
            _remove_paragraph(p)
            removed += 1
    return removed


def _sanitize_cell_paragraphs(cell) -> int:
    """Drop blue/※ guide runs; force remaining content to black."""
    removed = 0
    for p in cell.paragraphs:
        for run in list(p.runs):
            txt = run.text or ""
            if not txt.strip():
                continue
            drop = False
            if txt.strip().startswith("※"):
                drop = True
            elif _run_is_blue(run) and (
                INSTRUCTION_RUN_RE.search(txt)
                or txt.strip() in {"OOOOO", "OO.OO.OO", "OOO-OO-OOOOO"}
            ):
                drop = True
            if drop:
                run.text = ""
                removed += 1
                continue
            if _run_is_blue(run):
                _apply_black_font(run)
    return removed


def _sanitize_cell_deep(cell, stats: dict[str, int]) -> None:
    stats["cells_sanitized"] += 1
    stats["blue_runs_removed"] += _sanitize_cell_paragraphs(cell)
    for nested in cell.tables:
        for row in nested.rows:
            seen: set[int] = set()
            for ncell in row.cells:
                tc_id = id(ncell._tc)
                if tc_id in seen:
                    continue
                seen.add(tc_id)
                _sanitize_cell_deep(ncell, stats)


def _remove_duplicate_header_paragraphs(doc) -> int:
    """Remove template lines duplicated back-to-back (e.g. '1. 문제 인식…1. 문제 인식…')."""
    removed = 0
    for p in list(doc.paragraphs):
        t = (p.text or "").strip()
        if len(t) < 24:
            continue
        if len(t) % 2 == 0 and t[: len(t) // 2] == t[len(t) // 2 :]:
            _remove_paragraph(p)
            removed += 1
            continue
        m = re.match(r"^(.{20,}?)\1+$", t)
        if m:
            _remove_paragraph(p)
            removed += 1
    return removed


def _remove_instruction_paragraphs(doc) -> int:
    """Remove leftover guide lines; keep filled ㅇ answer paragraphs."""
    removed = 0
    for p in list(doc.paragraphs):
        t = (p.text or "").strip()
        if re.fullmatch(r"ㅇ\s*", t):
            _remove_paragraph(p)
            removed += 1
            continue
        if INSTRUCTION_PARA_RE.match(t) or t.startswith("※"):
            _remove_paragraph(p)
            removed += 1
    return removed


def _remove_paragraph(paragraph) -> None:
    el = paragraph._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def _sanitize_document(doc, *, pastes: dict[str, str] | None = None) -> dict[str, int]:
    stats: dict[str, Any] = {
        "blue_runs_removed": 0,
        "cells_sanitized": 0,
        "paragraphs_removed": 0,
        "body_flow_filled": {},
        "blue_body_paragraphs_removed": 0,
    }
    if pastes:
        stats["body_flow_filled"] = _fill_body_flow_answers(doc, pastes)
    stats["blue_body_paragraphs_removed"] = _purge_blue_instruction_body_paragraphs(doc)
    for table in doc.tables:
        for row in table.rows:
            seen_tc: set[int] = set()
            for cell in row.cells:
                tc_id = id(cell._tc)
                if tc_id in seen_tc:
                    continue
                seen_tc.add(tc_id)
                _sanitize_cell_deep(cell, stats)
    for p in list(doc.paragraphs):
        if _is_instruction_body_paragraph(p):
            _remove_paragraph(p)
            stats["blue_body_paragraphs_removed"] += 1
            continue
        for run in list(p.runs):
            txt = run.text or ""
            if _run_is_blue(run) and (
                txt.strip().startswith("※") or INSTRUCTION_RUN_RE.search(txt)
            ):
                run.text = ""
                stats["blue_runs_removed"] += 1
            elif _run_is_blue(run) and txt.strip():
                _apply_black_font(run)
    stats["paragraphs_removed"] = _remove_instruction_paragraphs(doc)
    stats["duplicate_header_paragraphs_removed"] = _remove_duplicate_header_paragraphs(doc)
    return stats


def _is_outline_toc_table(table) -> bool:
    pat = re.compile(r"^\d+\.\s+")
    hits = 0
    for row in table.rows:
        label = (row.cells[0].text or "").strip()
        if pat.match(label):
            hits += 1
    blob = _table_text(table, 30)
    if "기업명" in blob or "창업아이템명" in blob:
        return False
    return hits >= 3 and len(table.rows) <= 8


def _remove_boilerplate_tables(doc) -> int:
    removed = 0
    for table in list(doc.tables):
        text = _table_text(table, 40)
        if "※ 정부지원사업비" in text and len(table.rows) <= 2:
            table._element.getparent().remove(table._element)
            removed += 1
            continue
        if _is_outline_toc_table(table):
            table._element.getparent().remove(table._element)
            removed += 1
    return removed


def _table_text(table, max_rows: int = 30) -> str:
    parts: list[str] = []
    for row in table.rows[:max_rows]:
        for cell in row.cells:
            parts.append(cell.text or "")
    return "\n".join(parts)


def _table_has_marker(table, marker: str) -> bool:
    for row in table.rows[:8]:
        for cell in row.cells:
            if marker in (cell.text or ""):
                return True
    return False


def _find_table(doc, marker: str):
    for table in doc.tables:
        if _table_has_marker(table, marker):
            return table
    return None


def _remove_toc_tables(doc) -> int:
    removed = 0
    for table in list(doc.tables):
        text = "\n".join(
            (c.text or "") for row in table.rows[:6] for c in row.cells
        )
        if any(m in text for m in TOC_MARKERS) and "아이템 개요" not in text:
            table._element.getparent().remove(table._element)
            removed += 1
    removed += _remove_boilerplate_tables(doc)
    return removed


def _fill_company_table(table, entity: dict[str, Any]) -> dict[str, Any]:
    filled: dict[str, str] = {}
    name = entity.get("name") or ""
    brn = entity.get("brn") or ""
    founding = entity.get("founding_date") or "2021.01.05"
    address = entity.get("address") or "경기도 광명시 광명로 880"
    for row in table.rows:
        label = (row.cells[0].text or "").strip()
        if label == "기업명" and name and len(row.cells) > 1:
            _set_cell_text(row.cells[1], name)
            filled["company_name"] = name
            if len(row.cells) > 3:
                _set_cell_text(row.cells[3], founding)
                filled["founding_date"] = founding
        elif "사업자등록번호" in label and brn and len(row.cells) > 1:
            _set_cell_text(row.cells[1], brn)
            filled["brn"] = brn
            if len(row.cells) > 3:
                _set_cell_text(row.cells[3], address)
                filled["address"] = address
        elif "사업자 구분" in label and len(row.cells) > 1:
            _set_cell_text(row.cells[1], "법인사업자")
            filled["biz_type"] = "법인사업자"
        elif "대표자 유형" in label and len(row.cells) > 3:
            _set_cell_text(row.cells[3], "단독")
            filled["ceo_type"] = "단독"
    return filled


def _fill_general_status_table(table, item_line: str) -> dict[str, Any]:
    """창업아이템명·산출물·채용·분야·사업비·지역·팀·AI채용 요약 표."""
    out: dict[str, Any] = {}
    for row in table.rows:
        c0 = (row.cells[0].text or "").replace("\n", " ").strip()
        if c0 == "창업아이템명" and len(row.cells) > 3:
            _set_span(row, 3, len(row.cells), item_line)
            out["item_line"] = item_line
        elif "산출물" in c0 and len(row.cells) > 3:
            deliverables = (
                "모바일 어플리케이션(0개), 웹사이트(1), "
                "SaaS API 프로토타입(1), 사업계획서 PDF 파이프라인(1)"
            )
            _set_span(row, 3, len(row.cells), deliverables)
            out["deliverables"] = deliverables
        elif c0 == "채용인원(명)" and len(row.cells) > 3:
            _set_span(row, 3, 12, "1(명)")
            _set_span(row, 12, len(row.cells), "0(명)")
            out["hire_count"] = "정규직 1, 계약직 0"
        elif c0 == "지원 분야(택 1)":
            if len(row.cells) > 8:
                _check_cell(row.cells[8])
            out["support_field"] = "지식서비스"
        elif c0 == "전문기술분야(택 1)":
            for ci, cell in enumerate(row.cells):
                if "정보·통신" in (cell.text or "") and ci >= 2:
                    _check_cell(row.cells[ci - 1])
                    out["tech_field"] = "정보·통신"
                    break
        elif "총 사업비" in c0 and "00백만원" in (row.cells[3].text or ""):
            # col3-9 gov, 10-14 cash, 15-19 in-kind, 20-21 total (백만원)
            _set_span(row, 3, 10, "150백만원")
            _set_span(row, 10, 15, "20백만원")
            _set_span(row, 15, 20, "30백만원")
            _set_span(row, 20, len(row.cells), "200백만원")
            out["budget_summary"] = "150+20+30=200백만원"
        elif "지방우대 지역" in c0:
            if len(row.cells) > 13:
                _check_cell(row.cells[13])
            out["region"] = "일반지역"
        elif (
            c0 == "1"
            and "공동대표" in (row.cells[1].text if len(row.cells) > 1 else "")
            and "팀 구성 현황" in _table_text(table)
        ):
            _set_cell_text(row.cells[1], "대표")
            _set_span(row, 2, 6, "기획·RAG·무결성 게이트·사업 총괄")
            _set_span(row, 6, len(row.cells), TEAM_CAPABILITY_GENERAL)
            out["team_rep"] = True
        elif (
            c0 == "2"
            and "대리" in (row.cells[1].text if len(row.cells) > 1 else "")
            and "팀 구성 현황" in _table_text(table)
        ):
            _set_cell_text(row.cells[1], "AI인재(예정)")
            _set_span(row, 2, 6, "KB 인덱싱·슬롯 API·회귀 테스트")
            _set_span(
                row,
                6,
                len(row.cells),
                "중앙부처 AI 양성과정('23~'25) 수료 예정·Python/RAG",
            )
            out["team_ai_slot"] = True
        elif (
            c0 == "1"
            and (row.cells[1].text or "").strip() == "정규직"
            and "채용기간" in _table_text(table)
        ):
            _set_span(row, 2, 6, "선정 통보 D+60('26.08~)")
            _set_span(row, 6, 12, "40시간")
            _set_span(row, 12, len(row.cells), "RAG·슬롯 API·HOLD·감사 로그")
            out["ai_hire_row1"] = True
        elif c0 == "2" and "채용기간" in _table_text(table):
            for ci in range(1, len(row.cells)):
                _set_cell_text(row.cells[ci], "-")
            out["ai_hire_row2_cleared"] = True
    return out


def _fill_section_cover_table(table, item_line: str) -> dict[str, Any]:
    """사업계획 본문 표 상단 명칭·범주."""
    short = item_line.strip()[:90]
    scope = "B2G 정책·공공 서식 문서 초안 자동화(PoC, 자동제출·금융심사 대체 아님)"
    for row in table.rows:
        c0 = re.sub(r"\s+", "", (row.cells[0].text or ""))
        if c0 == "명칭" and len(row.cells) > 1:
            _set_span(row, 1, 3, short)
            if len(row.cells) > 4:
                _set_span(row, 4, len(row.cells), scope)
            return {"ok": True, "name": short, "scope": scope}
    return {"ok": False, "reason": "cover_name_row_not_found"}


def _fill_section_images(table) -> int:
    note = (
        "[파이프라인 요약]\n"
        "공고·서식 KB → 기업 프로필 질문 → RAG·슬롯 생성 → "
        "근거 미연결 HOLD → 담당자 확인 → PDF 초안\n"
        "(상세는 「실현 가능성」 본문 · 자동 제출 없음)"
    )
    filled = 0
    for row in table.rows:
        if (row.cells[0].text or "").strip() != "이미지" or len(row.cells) < 2:
            continue
        _set_section_cell_content(row.cells[1], note)
        written_tc = id(row.cells[1]._tc)
        for i in range(2, len(row.cells)):
            if id(row.cells[i]._tc) == written_tc:
                continue
            if (row.cells[i].text or "").strip():
                _clear_cell_content(row.cells[i])
                row.cells[i].add_paragraph()
        filled += 1
    return filled


def _fill_roadmap_dev_table(table) -> dict[str, Any]:
    rows_plan = [
        ("1", "AI 인재 채용", "2026.07 ~ 2026.09", "AI 양성과정 수료 인재 정규직 1명·4대보험"),
        (
            "2",
            "공고·서식 KB 구축",
            "2026.07 ~ 2026.10",
            "정책금융 공고·서식 수집·정규화·인덱싱 v0·감사 스키마",
        ),
        (
            "3",
            "RAG·PDF E2E 프로토타입",
            "2026.10 ~ 2026.12",
            "슬롯 생성 API·HOLD 플로우·골든셋 회귀 pytest",
        ),
        (
            "4",
            "hwp·운영 동결",
            "2027.01 ~ 2027.04",
            "hwp 필드 범위 확정·재인덱싱 절차·시연·B2B 파일럿 1건(LOI 시)",
        ),
    ]
    filled = 0
    ellipsis_row = None
    for row in table.rows:
        c0 = (row.cells[0].text or "").strip()
        for seq, title, period, detail in rows_plan:
            if c0 == seq:
                _set_cell_text(row.cells[1], title)
                _set_cell_text(row.cells[2], period)
                _set_cell_text(row.cells[3], detail)
                filled += 1
        if c0 == "…":
            ellipsis_row = row
    if ellipsis_row is not None and filled < len(rows_plan):
        seq, title, period, detail = rows_plan[filled]
        _set_cell_text(ellipsis_row.cells[0], seq)
        _set_cell_text(ellipsis_row.cells[1], title)
        _set_cell_text(ellipsis_row.cells[2], period)
        _set_cell_text(ellipsis_row.cells[3], detail)
        filled += 1
    elif ellipsis_row is not None:
        for cell in ellipsis_row.cells:
            _set_cell_text(cell, "")
    return {"rows": filled}


def _fill_roadmap_scale_table(table) -> dict[str, Any]:
    rows_plan = [
        ("1", "베타·내부 실증", "2026.10 ~ 2026.12", "골든셋 E2E·근거 연결율 내부 리포트"),
        ("2", "잠재고객 파일럿", "2027.01 ~ 2027.03", "관측 리포트 1건(면책·LOI 확보 시 명시)"),
        ("3", "B2B·API 탐색", "2027.02 ~ 2027.04", "구독·읽기전용 API 검토(협약 납품 범위 내)"),
        ("4", "성과·동결", "2027.04", "중간·최종 산출물·시연 스크립트"),
    ]
    filled = 0
    ellipsis_row = None
    for row in table.rows:
        c0 = (row.cells[0].text or "").strip()
        for seq, title, period, detail in rows_plan:
            if c0 == seq:
                _set_cell_text(row.cells[1], title)
                _set_cell_text(row.cells[2], period)
                _set_cell_text(row.cells[3], detail)
                filled += 1
        if c0 == "…":
            ellipsis_row = row
    if ellipsis_row is not None and filled < len(rows_plan):
        seq, title, period, detail = rows_plan[filled]
        _set_cell_text(ellipsis_row.cells[0], seq)
        _set_cell_text(ellipsis_row.cells[1], title)
        _set_cell_text(ellipsis_row.cells[2], period)
        _set_cell_text(ellipsis_row.cells[3], detail)
        filled += 1
    elif ellipsis_row is not None:
        for cell in ellipsis_row.cells:
            _set_cell_text(cell, "")
    return {"rows": filled}


def _fill_budget_table(table) -> dict[str, Any]:
    lines = [
        (
            "인건비",
            "▪ AI/RAG 엔지니어·창업자 인건비(12개월, 4대보험 포함)",
            "95,000,000",
            "10,000,000",
            "15,000,000",
            "120,000,000",
        ),
        (
            "지급수수료",
            "▪ 클라우드·벡터DB·LLM API·모니터링",
            "20,000,000",
            "",
            "",
            "20,000,000",
        ),
        (
            "외주용역비",
            "▪ UI/UX·보안 점검·변리 선행조사",
            "20,000,000",
            "5,000,000",
            "10,000,000",
            "35,000,000",
        ),
        (
            "소프트웨어",
            "▪ 개발도구·라이선스·테스트 데이터",
            "15,000,000",
            "5,000,000",
            "5,000,000",
            "25,000,000",
        ),
    ]
    totals = ("150,000,000", "20,000,000", "30,000,000", "200,000,000")

    def _apply_line(row, line: tuple[str, str, str, str, str, str]) -> None:
        _set_cell_text(row.cells[0], line[0])
        _set_cell_text(row.cells[1], line[1])
        for ci, val in enumerate(line[2:], start=2):
            if ci < len(row.cells):
                _set_cell_text(row.cells[ci], val)

    if len(table.rows) >= 3 + len(lines):
        data_rows = list(table.rows[3 : 3 + len(lines)])
    else:
        data_rows = [
            r
            for r in table.rows
            if (r.cells[0].text or "").replace("\t", "").strip()
            in ("재료비", "외주용역비", "…")
        ][: len(lines)]
    for row, line in zip(data_rows, lines):
        _apply_line(row, line)
    for row in table.rows:
        c0 = (row.cells[0].text or "").replace("\t", "").strip()
        if "합" in c0 and "계" in c0 and len(row.cells) >= 6:
            _set_cell_text(row.cells[2], totals[0])
            _set_cell_text(row.cells[3], totals[1])
            _set_cell_text(row.cells[4], totals[2])
            _set_cell_text(row.cells[5], totals[3])
    return {"line_items": len(lines), "total_krw": 200_000_000}


def _fill_team_detail_table(table, entity: dict[str, Any]) -> dict[str, Any]:
    del entity  # 법인 메타는 양식 필수칸만; 팀 표에는 성명·직장명 미기재(공고 PII 안내)
    for row in table.rows:
        c0 = (row.cells[0].text or "").strip()
        if c0 == "1":
            _set_cell_text(row.cells[1], "대표")
            _set_cell_text(row.cells[2], "기획·RAG·게이트·사업 총괄")
            _set_cell_text(row.cells[3], TEAM_CAPABILITY_DETAIL)
            if len(row.cells) > 4:
                _set_cell_text(row.cells[4], "완료('21.01)")
        elif c0 == "2":
            _set_cell_text(row.cells[1], "AI인재(예정)")
            _set_cell_text(row.cells[2], "KB·슬롯 API·회귀 테스트")
            _set_cell_text(
                row.cells[3],
                "AI 양성과정('23~'25) 수료·Python·벡터검색/RAG",
            )
            if len(row.cells) > 4:
                _set_cell_text(row.cells[4], "예정('26.08)")
        elif c0 == "…":
            _set_cell_text(row.cells[0], "3")
            _set_cell_text(row.cells[1], "-")
            _set_cell_text(row.cells[2], "-")
            _set_cell_text(row.cells[3], "-")
            if len(row.cells) > 4:
                _set_cell_text(row.cells[4], "-")
    return {"ok": True}


def _fill_partner_table(table) -> dict[str, Any]:
    for row in table.rows:
        c0 = (row.cells[0].text or "").strip()
        if c0 == "1":
            _set_cell_text(row.cells[1], "자체 개발(1인+채용)")
            _set_cell_text(row.cells[2], "RAG·게이트·문서 파이프라인")
            _set_cell_text(row.cells[3], "핵심 개발 자체 수행, 외주는 UI·보안 등 제한적")
            _set_cell_text(row.cells[4], "2026.07~")
        elif c0 == "2":
            _set_cell_text(row.cells[1], "K-Startup·정책금융 공고 채널")
            _set_cell_text(row.cells[2], "공고·서식·버전 메타데이터")
            _set_cell_text(row.cells[3], "KB 수집·재인덱싱·회귀 테스트 입력")
            _set_cell_text(row.cells[4], "2026.07~")
        elif c0 in ("...", "…"):
            for cell in row.cells:
                _set_cell_text(cell, "")
    return {"ok": True}


def _fill_ai_schedule_table(table) -> dict[str, Any]:
    for row in table.rows:
        c0 = (row.cells[0].text or "").strip()
        if c0 == "1":
            _set_cell_text(row.cells[1], "정규직")
            _set_cell_text(row.cells[2], "RAG·슬롯 API·HOLD·감사 로그")
            _set_cell_text(row.cells[3], "40시간")
            if len(row.cells) > 4:
                _set_span(row, 4, len(row.cells) - 1, "RAG 백엔드·회귀 pytest")
            if len(row.cells) > 5:
                _set_cell_text(row.cells[-1], "예정('26.08)")
        elif c0 == "2":
            _set_cell_text(row.cells[1], "-")
            _set_cell_text(row.cells[2], "-")
            _set_cell_text(row.cells[3], "-")
            for i in range(4, len(row.cells)):
                _set_cell_text(row.cells[i], "-")
        elif c0 == "…":
            for cell in row.cells:
                _set_cell_text(cell, "")
    return {"ok": True}


def _fill_milestone_table(table) -> dict[str, Any]:
    patches = {
        "...": ("월별 점검", "'27.2", "골든셋 회귀·근거 연결율 주간 리포트"),
        "5": None,
    }
    detail_rows = {
        5: ("제품 고도화", "'27.3", "hwp 필드 확장·재인덱싱·파일럿 보완"),
        6: ("최종 납품", "'27.4", "협약 산출물·시연·감사 로그 패키지"),
    }
    for ri, row in enumerate(table.rows):
        c0 = (row.cells[0].text or "").strip()
        if c0 in patches and patches[c0]:
            _set_cell_text(row.cells[0], patches[c0][0])
            _set_cell_text(row.cells[1], patches[c0][1])
            _set_cell_text(row.cells[2], patches[c0][2])
        if ri in detail_rows:
            _set_cell_text(row.cells[0], detail_rows[ri][0])
            _set_cell_text(row.cells[1], detail_rows[ri][1])
            _set_cell_text(row.cells[2], detail_rows[ri][2])
    return {"ok": True}


def _find_table_by_text(
    doc,
    needle: str,
    *,
    also: str | None = None,
    min_cols: int = 0,
    max_cols: int = 99,
    must_have: str | None = None,
):
    for table in doc.tables:
        if not table.rows:
            continue
        ncol = len(table.rows[0].cells)
        if ncol < min_cols or ncol > max_cols:
            continue
        blob = _table_text(table)
        if needle not in blob:
            continue
        if also is not None and also not in blob:
            continue
        if must_have is not None and must_have not in blob:
            continue
        return table
    return None


def _fill_item_name_row(table, item_line: str) -> bool:
    for row in table.rows:
        if "창업아이템명" not in (row.cells[0].text or ""):
            continue
        if len(row.cells) < 4:
            return False
        _set_cell_text(row.cells[3], item_line)
        return True
    return False


def _enrich_section_paste(marker: str, fname: str, text: str, pastes: dict[str, str]) -> str:
    """양식 소제목(검정) + paste — 중첩 ※ 안내 칸 대체."""
    if marker == "아이템 개요":
        return (
            "■ 창업 아이템 개요(요약)\n\n"
            f"{text}\n\n"
            "■ 개발 아이템 소개\n\n"
            "공공·정책 신청 문서를 공고·서식·근거와 연동해 초안 생성하는 B2G 문서 자동화 PoC "
            "(자동 제출·금융 심사 대체 아님)."
        )
    if marker == "문제 인식":
        return (
            "■ 창업 아이템의 필요성 (국내·외 시장 현황 및 문제점)\n\n"
            f"{text}\n\n"
            "■ 문제 해결을 위한 개발 필요성\n\n"
            "근거·서식·버전이 결합된 라스트마일 문서 파이프라인 없이는 "
            "공공 신청 품질·감사 추적을 안정적으로 확보하기 어렵다."
        )
    if marker == "실현 가능성":
        p04 = pastes.get("plan_04_growth_funding_paste.txt", "")
        budget_note = ""
        if p04:
            m = re.search(r"4\.3[^\n]*\n+(.*?)(?:\n\n|$)", p04, re.DOTALL)
            if m:
                budget_note = m.group(1).strip()[:400]
        extra = (
            f"\n\n■ 사업비 집행 연계\n\n{budget_note}"
            if budget_note
            else "\n\n■ 사업비 집행 연계\n\n집행 상세는 본문 「사업비 집행 계획」표 및 일반현황 총사업비 표 참조."
        )
        return (
            "■ 창업 아이템의 개발 계획 (제품·서비스 구체화)\n\n"
            f"{text}{extra}"
        )
    return text


def _fill_section_table(table, pastes: dict[str, str]) -> dict[str, Any]:
    sections: dict[str, Any] = {}
    for row in table.rows:
        label = (row.cells[0].text or "").replace("\n", " ").strip()
        if len(row.cells) < 2:
            continue
        for marker, fname in SECTION_MAP:
            if marker not in label:
                continue
            text = pastes.get(fname, "")
            if not text:
                sections[fname] = {"ok": False, "reason": "paste_missing"}
                break
            if marker == "AI 인재 활용":
                text = ai_talent_body_summary(text)
            else:
                text = _enrich_section_paste(marker, fname, text, pastes)
            n = _set_section_cell_content(row.cells[1], text)
            sections[fname] = {"ok": n >= 80, "chars": n, "label": label}
            break
    return sections


def _extract_item_title(pastes: dict[str, str]) -> str:
    body = pastes.get("plan_01_summary_paste.txt", "")
    m = re.search(r"과제명\(안\):\s*(.+)", body)
    if m:
        return m.group(1).strip()
    first = body.splitlines()[0].strip("- ").strip() if body else ""
    return first or "공공·정책 신청 문서 AI 초안 생성 플랫폼"


def _export_hwp_hancom(docx_path: Path, hwp_path: Path) -> dict[str, Any]:
    try:
        from pyhwpx import Hwp  # noqa: PLC0415
    except ImportError:
        return {"ok": False, "reason": "pyhwpx_missing"}
    hwp = None
    try:
        hwp_path.parent.mkdir(parents=True, exist_ok=True)
        if hwp_path.exists():
            hwp_path.unlink()
        hwp = Hwp(visible=False)
        hwp.open(str(docx_path.resolve()))
        hwp.save_as(str(hwp_path.resolve()), format="HWP")
        hwp.quit()
        hwp = None
        size = hwp_path.stat().st_size if hwp_path.is_file() else 0
        return {"ok": hwp_path.is_file() and size > 1024, "hwp": str(hwp_path), "bytes": size}
    except Exception as exc:
        if hwp is not None:
            try:
                hwp.quit()
            except Exception:
                pass
        return {"ok": False, "error": str(exc)}


def _export_pdf_word(docx_path: Path, pdf_path: Path) -> dict[str, Any]:
    try:
        import win32com.client  # noqa: PLC0415
    except ImportError:
        return {"ok": False, "reason": "win32com_missing"}
    word = None
    try:
        pdf_path.parent.mkdir(parents=True, exist_ok=True)
        if pdf_path.exists():
            pdf_path.unlink()
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(str(docx_path.resolve()))
        doc.ExportAsFixedFormat(str(pdf_path.resolve()), ExportFormat=17)
        doc.Close(False)
        word.Quit()
        word = None
        size = pdf_path.stat().st_size if pdf_path.is_file() else 0
        return {"ok": pdf_path.is_file() and size > 1024, "pdf": str(pdf_path), "bytes": size}
    except Exception as exc:
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass
        return {"ok": False, "error": str(exc)}


def _export_ai_talent_pdf(text: str, pdf_path: Path) -> dict[str, Any]:
    from docx import Document

    import tempfile

    tmp = Path(tempfile.gettempdir()) / f"mkm_ai_talent_2p_{datetime.now().strftime('%Y%m%d%H%M%S')}.docx"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    d = Document()
    d.add_heading("AI 인재 활용 계획 (2p 요약)", level=1)
    for para in text.split("\n"):
        d.add_paragraph(para)
    d.save(str(tmp))
    out = _export_pdf_word(tmp, pdf_path)
    out["via"] = "scratch_docx"
    return out


def fill_docx(*, template: Path, out_docx: Path) -> dict[str, Any]:
    from docx import Document

    cfg = _load_config()
    elig = _load_eligibility()
    entity = _entity_profile(cfg, elig)
    pastes = _load_pastes()
    if len(pastes) < 5:
        return {"ok": False, "error": "paste_pack_incomplete", "pastes": list(pastes)}

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, out_docx)
    doc = Document(str(out_docx))

    removed = _remove_toc_tables(doc)
    company_t = _find_table(doc, "기업명")
    general_t = _find_table(doc, "창업아이템명")
    section_t = _find_table(doc, "아이템 개요")
    dev_roadmap_t = _find_table_by_text(
        doc,
        "AI 인재 채용",
        also="추진 기간",
        min_cols=4,
        max_cols=5,
        must_have="프로토타입",
    )
    budget_t = _find_table_by_text(
        doc,
        "집행 계획",
        min_cols=6,
        max_cols=6,
        must_have="재료비",
    )
    scale_roadmap_t = _find_table_by_text(
        doc, "시제품 설계", also="추진 기간", min_cols=4, max_cols=5, must_have="00년 상반기"
    )
    team_t = _find_table_by_text(
        doc,
        "보유 역량(경력 및 학력 등)",
        also="구성 상태",
        min_cols=5,
        max_cols=5,
        must_have="공동대표",
    )
    partner_t = _find_table_by_text(
        doc, "파트너명", also="협업 방안", min_cols=5, max_cols=5, must_have="○○전자"
    )
    ai_sched_t = _find_table_by_text(
        doc,
        "채용유형",
        also="구성 상태",
        min_cols=5,
        max_cols=6,
        must_have="AX 로드맵",
    )
    milestone_t = _find_table_by_text(
        doc, "핵심 개발과제 추진", also="4대보험", min_cols=3, max_cols=3
    )

    meta: dict[str, Any] = {
        "ok": False,
        "removed_toc_tables": removed,
        "tables_found": {
            "company": company_t is not None,
            "general": general_t is not None,
            "sections": section_t is not None,
            "dev_roadmap": dev_roadmap_t is not None,
            "budget": budget_t is not None,
            "scale_roadmap": scale_roadmap_t is not None,
            "team": team_t is not None,
            "partner": partner_t is not None,
            "ai_schedule": ai_sched_t is not None,
            "milestone": milestone_t is not None,
        },
    }
    if not section_t:
        meta["error"] = "section_table_not_found"
        return meta

    item = _extract_item_title(pastes)
    if company_t:
        meta["company"] = _fill_company_table(company_t, entity)
    if general_t:
        meta["item_name"] = {"text": item, "ok": _fill_item_name_row(general_t, item)}
        meta["general_status"] = _fill_general_status_table(general_t, item)
    meta["section_cover"] = _fill_section_cover_table(section_t, item)
    meta["sections"] = _fill_section_table(section_t, pastes)
    meta["section_images"] = _fill_section_images(section_t)
    if dev_roadmap_t:
        meta["dev_roadmap"] = _fill_roadmap_dev_table(dev_roadmap_t)
    if budget_t:
        meta["budget"] = _fill_budget_table(budget_t)
    if scale_roadmap_t:
        meta["scale_roadmap"] = _fill_roadmap_scale_table(scale_roadmap_t)
    if team_t:
        meta["team_detail"] = _fill_team_detail_table(team_t, entity)
    if partner_t:
        meta["partners"] = _fill_partner_table(partner_t)
    if ai_sched_t and ai_sched_t is not team_t:
        meta["ai_schedule"] = _fill_ai_schedule_table(ai_sched_t)
    if milestone_t:
        meta["milestones"] = _fill_milestone_table(milestone_t)

    meta["sanitize"] = _sanitize_document(doc, pastes=pastes)

    doc.save(str(out_docx))
    ok_sections = sum(
        1 for v in meta["sections"].values() if isinstance(v, dict) and v.get("ok")
    )
    meta["sections_ok_count"] = ok_sections
    cover_ok = bool((meta.get("section_cover") or {}).get("ok"))
    meta["ok"] = ok_sections >= 5 and cover_ok
    return meta


def main() -> int:
    ap = argparse.ArgumentParser(description="Fill 도약 별첨1 from paste_ready.")
    ap.add_argument("--template", default="", help="Override template docx path")
    ap.add_argument("--out-docx", default=str(OUT_DOCX))
    ap.add_argument("--out-pdf", default=str(OUT_PDF))
    ap.add_argument("--out-ai-pdf", default=str(OUT_AI_PDF))
    ap.add_argument("--skip-pdf", action="store_true")
    ap.add_argument("--skip-ai-pdf", action="store_true")
    ap.add_argument("--out-hwp", default="", help="Optional Hancom HWP export path (pyhwpx)")
    ap.add_argument("--skip-hwp", action="store_true")
    args = ap.parse_args()

    cfg = _load_config()
    att = cfg.get("attachments") or {}
    template = Path(args.template) if args.template else ROOT / str(
        att.get("doyak_plan_docx") or ""
    )
    if not template.is_file():
        print(f"template missing: {template}", flush=True)
        return 1

    out_docx = Path(args.out_docx).resolve()
    out_pdf = Path(args.out_pdf).resolve()
    out_ai_pdf = Path(args.out_ai_pdf).resolve()

    fill_meta = fill_docx(template=template, out_docx=out_docx)
    pastes = _load_pastes()

    pdf_meta: dict[str, Any] = {"skipped": True}
    if not args.skip_pdf:
        pdf_meta = _export_pdf_word(out_docx, out_pdf)

    ai_meta: dict[str, Any] = {"skipped": True}
    if not args.skip_ai_pdf:
        ai_text = pastes.get("plan_06_ai_talent_2p_paste.txt", "")
        if ai_text:
            ai_meta = _export_ai_talent_pdf(ai_text, out_ai_pdf)

    hwp_meta: dict[str, Any] = {"skipped": True}
    out_hwp: Path | None = None
    if not args.skip_hwp:
        out_hwp = Path(args.out_hwp).resolve() if args.out_hwp else out_docx.with_suffix(".hwp")
        hwp_meta = _export_hwp_hancom(out_docx, out_hwp)

    report = {
        "schema": "kstartup_startup_package_ai_filled_plan_v1",
        "generated_at_utc": _utc(),
        "template": str(template.relative_to(ROOT)).replace("\\", "/"),
        "outputs": {
            "docx": str(out_docx.relative_to(ROOT)).replace("\\", "/"),
            "pdf": str(out_pdf.relative_to(ROOT)).replace("\\", "/"),
            "ai_talent_pdf": str(out_ai_pdf.relative_to(ROOT)).replace("\\", "/"),
            "hwp": str(out_hwp.relative_to(ROOT)).replace("\\", "/") if out_hwp else "",
        },
        "fill": fill_meta,
        "pdf_export": pdf_meta,
        "ai_talent_pdf_export": ai_meta,
        "hwp_export": hwp_meta,
        "ok": bool(fill_meta.get("ok") and (pdf_meta.get("ok") or args.skip_pdf)),
    }
    OUT_META.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
