# -*- coding: utf-8 -*-
"""[HYPO] Export Han Vocology MD chapters + embedded PNG figures to DOCX.

Manifest: docs/final/artifacts/han_vocology_docx_export_manifest_v1.json
Requires: python-docx (pip install python-docx)
"""

from __future__ import annotations

import argparse
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_MANIFEST = ROOT / "docs/final/artifacts/han_vocology_docx_export_manifest_v1.json"
IMG_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")


def load_manifest(path: Path) -> dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if data.get("schema") != "han_vocology_docx_export_manifest_v1":
        raise ValueError(f"unexpected manifest schema: {data.get('schema')}")
    return data


def resolve_md_path(md_path: str, chapter_file: Path) -> Path:
    p = Path(md_path)
    if p.is_absolute():
        return p
    if p.parts and p.parts[0] == "..":
        return (chapter_file.parent / p).resolve()
    candidate = ROOT / p
    if candidate.is_file():
        return candidate
    return (chapter_file.parent / p).resolve()


def strip_inline_md(text: str) -> str:
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()


def add_markdown_to_doc(doc: Any, md_text: str, chapter_file: Path) -> int:
    from docx.shared import Inches

    images = 0
    for raw_line in md_text.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            continue
        if line.strip() == "---":
            doc.add_paragraph("")
            continue
        img = IMG_RE.search(line)
        if img and line.strip().startswith("!"):
            img_path = resolve_md_path(img.group(2), chapter_file)
            if img_path.is_file():
                doc.add_picture(str(img_path), width=Inches(5.8))
                cap = img.group(1).strip()
                if cap:
                    doc.add_paragraph(cap)
                images += 1
            else:
                doc.add_paragraph(f"[이미지 없음: {img_path}]")
            continue
        if line.startswith("# "):
            doc.add_heading(strip_inline_md(line[2:]), level=1)
        elif line.startswith("## "):
            doc.add_heading(strip_inline_md(line[3:]), level=2)
        elif line.startswith("### "):
            doc.add_heading(strip_inline_md(line[4:]), level=3)
        elif line.startswith("> "):
            p = doc.add_paragraph(strip_inline_md(line[2:]))
            p.style = "Intense Quote"
        elif line.startswith("|"):
            doc.add_paragraph(strip_inline_md(line))
        elif line.startswith("- "):
            doc.add_paragraph(strip_inline_md(line[2:]), style="List Bullet")
        elif re.match(r"^\d+\.\s", line):
            doc.add_paragraph(strip_inline_md(re.sub(r"^\d+\.\s", "", line)), style="List Number")
        else:
            doc.add_paragraph(strip_inline_md(line))
    return images


def export_volume(
    volume: dict[str, Any],
    out_path: Path,
    *,
    workspace_root: Path = ROOT,
) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError as exc:
        raise SystemExit("python-docx required: pip install python-docx") from exc

    doc = Document()
    doc.add_heading(volume.get("title_ko", "한의음성학"), level=0)
    disclaimer = (
        "[B-track · 교육 초안 · 진단·처방 대체 아님] "
        f"Generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}. "
        "도판 schematic=internal_original; F12=licensed base+overlay (adjudication pending). "
        "send_gate: HOLD."
    )
    doc.add_paragraph(disclaimer)

    chapters_included: list[str] = []
    missing: list[str] = []
    image_count = 0

    for rel in volume.get("chapters", []):
        chapter_path = workspace_root / rel
        if not chapter_path.is_file():
            missing.append(rel)
            continue
        doc.add_page_break()
        doc.add_heading(chapter_path.stem, level=1)
        md_text = chapter_path.read_text(encoding="utf-8")
        image_count += add_markdown_to_doc(doc, md_text, chapter_path)
        doc.add_paragraph(
            "— 라이선스: 본 장 도판은 han_vocology_figure_registry_v1_latest.json 및 "
            "HAN_VOCOLOGY_FIGURE_LICENSE_MAP_V0_1.md 참조. 외부 배포 전 send_gate 확인."
        )
        chapters_included.append(rel)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)

    return {
        "ok": True,
        "schema": "han_vocology_docx_export_report_v1",
        "output": out_path.as_posix(),
        "title_ko": volume.get("title_ko"),
        "chapters_included": chapters_included,
        "chapters_missing": missing,
        "image_count": image_count,
        "chapter_count": len(chapters_included),
    }


def main() -> int:
    ap = argparse.ArgumentParser(description="Export Han Vocology MD bundle to DOCX")
    ap.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    ap.add_argument("--volume", default="pilot", help="pilot | volume_a_core | volume_b")
    ap.add_argument("--out", type=Path, default=None)
    ap.add_argument("--report-out", type=Path, default=None)
    args = ap.parse_args()

    manifest = load_manifest(args.manifest)
    volumes = manifest.get("volumes") or {}
    if args.volume not in volumes:
        raise SystemExit(f"unknown volume: {args.volume}; keys={list(volumes)}")

    volume = volumes[args.volume]
    out_path = args.out or (ROOT / volume.get("default_out", f"reports/han_vocology_export_{args.volume}_latest.docx"))
    report = export_volume(volume, out_path)
    report["generated_at_utc"] = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    report["volume_key"] = args.volume

    report_path = args.report_out or out_path.with_suffix(".json")
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"ok": True, "output": report["output"], "chapters": report["chapter_count"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
