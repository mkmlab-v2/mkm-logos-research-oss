#!/usr/bin/env python3
from __future__ import annotations

import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
FILLED = ROOT / "reports/kstartup_startup_package_ai_filled/doyak_plan_filled_v5.docx"
TEMPLATE = ROOT / (
    "reports/kstartup_startup_package_ai_attachments/"
    "(도약) [별첨 1] 2026년 창업도약패키지(AI 인재 실증형) 사업계획서 양식.docx"
)
BLUE_THEME = re.compile(r'w:color w:val="([0-9A-Fa-f]{6})"')
BLUE_ATTR = re.compile(r'w:val="(0000FF|0070C0|4472C4|0563C1|5B9BD5|365F91|2E75B6)"')


def xml_blue_snippets(path: Path, limit: int = 12) -> list[str]:
    xml = ZipFile(path).read("word/document.xml").decode("utf-8")
    out: list[str] = []
    for m in re.finditer(r"<w:t[^>]*>([^<]{3,120})</w:t>", xml):
        start = max(0, m.start() - 400)
        ctx = xml[start : m.end()]
        if BLUE_ATTR.search(ctx) or "themeColor" in ctx:
            if "※" in m.group(1) or "개발하고자" in m.group(1) or "ㅇ" in m.group(1):
                out.append(m.group(1))
        if len(out) >= limit:
            break
    return out


def scan_doc(path: Path, label: str) -> None:
    print(f"\n===== {label} =====")
    doc = Document(str(path))
    for ti, table in enumerate(doc.tables):
        blob = "".join(c.text for r in table.rows for c in r.cells)
        if "아이템 개요" not in blob:
            continue
        print(f"section_table T{ti} rows={len(table.rows)}")
        for ri, row in enumerate(table.rows):
            c0 = (row.cells[0].text or "").replace("\n", " | ")
            c1 = row.cells[1].text if len(row.cells) > 1 else ""
            if ri <= 8 or "문제" in c0:
                inst = "※" in c1 or "개발하고자" in c1
                print(
                    f" r{ri:02d} c0={c0[:50]!r} c1_len={len(c1)} "
                    f"nested={len(row.cells[1].tables) if len(row.cells)>1 else 0} inst={inst}"
                )
                if ri == 2:
                    print(f"      BODY={c1[:300]!r}")
    print("--- body paragraphs (problem/blue) ---")
    for i, p in enumerate(doc.paragraphs):
        t = (p.text or "").strip()
        if not t:
            continue
        if any(
            k in t
            for k in (
                "문제 인식",
                "창업 아이템의 필요성",
                "개발 아이템",
                "개발하고자",
                "※",
                "Problem",
            )
        ):
            print(f" P{i:03d}: {t[:140]!r}")
    print("--- xml blue instruction snippets ---")
    for s in xml_blue_snippets(path):
        print(" ", repr(s[:100]))


if __name__ == "__main__":
    if FILLED.is_file():
        scan_doc(FILLED, "FILLED v5")
    if TEMPLATE.is_file():
        scan_doc(TEMPLATE, "TEMPLATE")
