#!/usr/bin/env python3
"""Verify filled 도약 별첨1 docx/PDF has real paste content (G1 pre-attest)."""
from __future__ import annotations

import argparse
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]

from kstartup_startup_package_ai_form_notice_compliance_v1 import (  # noqa: E402
    AI_TALENT_BODY_SUMMARY_MAX_CHARS,
    compliance_report_from_docx_tables,
)
META = ROOT / "reports/kstartup_startup_package_ai_filled_plan_latest.json"
DEFAULT_DOCX = ROOT / "reports/kstartup_startup_package_ai_filled/doyak_plan_filled_v1.docx"
DEFAULT_PDF = ROOT / "reports/kstartup_startup_package_ai_filled/doyak_plan_filled_v1.pdf"
OUT = ROOT / "reports/kstartup_startup_package_ai_filled_plan_verify_latest.json"

MIN_SECTION_CHARS = 80
SECTION_MIN_OVERRIDES: dict[str, int] = {
    "아이템 개요": 280,
    "성장전략": 280,
    "팀 구성": 180,
}
MIN_PDF_BYTES = 20_000
BODY_PAGES_TARGET = 15
BODY_PAGES_HARD_MAX = 18
AI_SECTION_PAGES_MAX = 2
SECTION_LABELS = (
    "아이템 개요",
    "문제 인식",
    "실현 가능성",
    "성장전략",
    "팀 구성",
    "AI 인재 활용",
)

PLACEHOLDER_PATTERNS: tuple[str | re.Pattern[str], ...] = (
    "OO.OO.OO",
    "OO도 OO시",
    re.compile(r"(?<![0-9A-Za-z가-힣])OOOOO(?![0-9A-Za-z가-힣])"),
    "00(명)",
    re.compile(r"(?<![0-9])00백만원"),
    "00.00 ~ 00.00",
    "00년 상반기",
    "OO학",
    "○○전자",
    "< 사진",
    "DMD소켓",
    "시금형제작",
)


def _cell_has_placeholder(tx: str, pat: str | re.Pattern[str]) -> bool:
    if isinstance(pat, re.Pattern):
        return bool(pat.search(tx))
    return pat in tx


def _utc() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _placeholder_hits_from_docx(docx: Path) -> list[str]:
    from docx import Document

    hits: list[str] = []
    doc = Document(str(docx))
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                tx = (cell.text or "").strip()
                if not tx:
                    continue
                for pat in PLACEHOLDER_PATTERNS:
                    if _cell_has_placeholder(tx, pat):
                        label = pat.pattern if isinstance(pat, re.Pattern) else pat
                        hits.append(f"T{ti}r{ri}c{ci}:{label}")
                        break
    return hits


def _section_cover_name_from_docx(docx: Path) -> str:
    from docx import Document

    doc = Document(str(docx))
    for table in doc.tables:
        for row in table.rows:
            c0 = re.sub(r"\s+", "", (row.cells[0].text or ""))
            if c0 == "명칭" and len(row.cells) > 1:
                return (row.cells[1].text or "").strip()
    return ""


def _ai_hire_row2_residual_hits(docx: Path) -> list[str]:
    from docx import Document

    doc = Document(str(docx))
    hits: list[str] = []
    for ti, table in enumerate(doc.tables):
        blob = "".join(c.text for r in table.rows for c in r.cells)
        if "채용기간" not in blob or "AI 인재 채용 계획" not in blob:
            continue
        for ri, row in enumerate(table.rows):
            c0 = (row.cells[0].text or "").strip()
            if c0 != "2":
                continue
            row_text = " ".join((c.text or "") for c in row.cells)
            if "ROI" in row_text or "성과측정" in row_text:
                hits.append(f"T{ti}r{ri}:roi_residual")
            if re.search(r"['\u2019]26\.[0-9]", row_text) and "-" not in row_text[:40]:
                hits.append(f"T{ti}r{ri}:hire_date_residual")
    return hits


def _section_chars_from_docx(docx: Path) -> dict[str, int]:
    from docx import Document

    doc = Document(str(docx))
    counts: dict[str, int] = {}
    for table in doc.tables:
        for row in table.rows:
            label = (row.cells[0].text or "").replace("\n", " ").strip()
            if len(row.cells) < 2:
                continue
            for marker in SECTION_LABELS:
                if marker in label:
                    counts[marker] = len((row.cells[1].text or "").strip())
    return counts


def _word_page_stats(docx: Path) -> dict[str, Any]:
    try:
        import win32com.client  # noqa: PLC0415
    except ImportError:
        return {"ok": False, "reason": "win32com_missing"}
    word = None
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(str(docx.resolve()))
        doc.Repaginate()
        total = int(doc.Range().Information(4))
        doc.Close(False)
        word.Quit()
        word = None
        return {"ok": True, "total_pages": total}
    except Exception as exc:
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass
        return {"ok": False, "reason": str(exc)}


def _paragraph_xml_text(paragraph) -> str:
    return "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", paragraph._element.xml))


def _nested_instruction_hits(docx: Path) -> list[str]:
    from docx import Document

    doc = Document(str(docx))
    hits: list[str] = []
    for ti, table in enumerate(doc.tables):
        if "아이템 개요" not in "".join(c.text for r in table.rows for c in r.cells):
            continue
        for ri, row in enumerate(table.rows):
            if len(row.cells) < 2:
                continue
            cell = row.cells[1]
            for nti, nested in enumerate(cell.tables):
                for nri, nrow in enumerate(nested.rows):
                    for nci, ncell in enumerate(nrow.cells):
                        tx = ncell.text or ""
                        if "※" in tx or "개발하고자" in tx:
                            hits.append(f"T{ti}r{ri}n{nti}r{nri}c{nci}")
    for i, p in enumerate(doc.paragraphs):
        xt = _paragraph_xml_text(p)
        t = p.text or ""
        if xt.strip().startswith("※") or t.strip().startswith("※"):
            hits.append(f"P{i}:※")
        elif "개발하고자" in xt and "0000FF" in p._element.xml.upper():
            hits.append(f"P{i}:blue_guide")
        elif re.fullmatch(r"ㅇ\s*", t.strip()):
            hits.append(f"P{i}:empty_o")
        elif t.strip().startswith("ㅇ") and "세부내용" in t:
            hits.append(f"P{i}:guide_o")
    return hits


BLUE_RGB = frozenset(
    {"0000FF", "0070C0", "4472C4", "2F5496", "0563C1", "5B9BD5", "2E75B6"}
)


def _blue_run_count(docx: Path) -> int:
    from docx import Document

    doc = Document(str(docx))
    n = 0
    for table in doc.tables:
        for row in table.rows:
            seen: set[int] = set()
            for cell in row.cells:
                tc = id(cell._tc)
                if tc in seen:
                    continue
                seen.add(tc)
                for p in cell.paragraphs:
                    for run in p.runs:
                        col = run.font.color
                        if col is not None and col.rgb is not None:
                            if str(col.rgb).upper() in BLUE_RGB and (run.text or "").strip():
                                n += 1
    for p in doc.paragraphs:
        xml = p._element.xml.upper()
        if any(rgb in xml for rgb in BLUE_RGB):
            xt = _paragraph_xml_text(p).strip()
            if xt and ("※" in xt or "개발하고자" in xt or "아이디어를" in xt):
                n += 1
    return n


def verify(*, docx: Path, pdf: Path, min_section: int, min_pdf: int) -> dict[str, Any]:
    reasons: list[str] = []
    section_chars: dict[str, int] = {}

    placeholder_hits: list[str] = []
    nested_instruction_hits: list[str] = []
    blue_runs = 0
    form_notice: dict[str, Any] = {}
    page_stats: dict[str, Any] = {}
    if not docx.is_file():
        reasons.append("filled_docx_missing")
    else:
        from docx import Document

        section_chars = _section_chars_from_docx(docx)
        placeholder_hits = _placeholder_hits_from_docx(docx)
        blue_runs = _blue_run_count(docx)
        form_notice = compliance_report_from_docx_tables(Document(str(docx)))
        nested_instruction_hits = _nested_instruction_hits(docx)
        page_stats = _word_page_stats(docx)
        if blue_runs > 0:
            reasons.append(f"blue_runs_remain:{blue_runs}")
        if nested_instruction_hits:
            reasons.append(f"nested_instruction_remain:{len(nested_instruction_hits)}")
        if form_notice.get("pii_hits"):
            reasons.append(f"pii_hits:{len(form_notice['pii_hits'])}")
        if form_notice.get("instruction_hits"):
            reasons.append(f"instruction_remnants:{len(form_notice['instruction_hits'])}")
        ai_chars = section_chars.get("AI 인재 활용", 0)
        if ai_chars > AI_TALENT_BODY_SUMMARY_MAX_CHARS:
            reasons.append(f"ai_section_chars_over_2p_budget:{ai_chars}")
        if page_stats.get("ok"):
            total_p = int(page_stats.get("total_pages") or 0)
            if total_p > BODY_PAGES_HARD_MAX:
                reasons.append(f"body_pages_over_hard_max:{total_p}")
            ai_p = page_stats.get("ai_section_pages")
            if ai_p is not None and int(ai_p) > AI_SECTION_PAGES_MAX:
                reasons.append(f"ai_section_pages_over:{ai_p}")
        thin = [
            k
            for k, n in section_chars.items()
            if n < SECTION_MIN_OVERRIDES.get(k, min_section)
        ]
        if len(section_chars) < 5:
            reasons.append("section_rows_lt_5")
        if thin:
            reasons.append(f"thin_sections:{','.join(thin)}")
        if placeholder_hits:
            reasons.append(f"placeholders_remain:{len(placeholder_hits)}")
        cover_name = _section_cover_name_from_docx(docx)
        if not cover_name or cover_name in ("명칭", "명     칭"):
            reasons.append("section_cover_name_empty")
        ai_row2 = _ai_hire_row2_residual_hits(docx)
        if ai_row2:
            reasons.append(f"ai_hire_row2_residual:{','.join(ai_row2)}")

    pdf_bytes = 0
    if not pdf.is_file():
        reasons.append("filled_pdf_missing")
    else:
        pdf_bytes = pdf.stat().st_size
        if pdf_bytes < min_pdf:
            reasons.append(f"pdf_too_small:{pdf_bytes}")

    return {
        "schema": "kstartup_startup_package_ai_filled_plan_verify_v1",
        "generated_at_utc": _utc(),
        "docx": str(docx),
        "pdf": str(pdf),
        "section_chars": section_chars,
        "placeholder_hits": placeholder_hits,
        "nested_instruction_hits": nested_instruction_hits,
        "blue_runs": blue_runs,
        "form_notice_compliance": form_notice,
        "page_stats": page_stats,
        "form_notice_policy": {
            "body_pages_target": BODY_PAGES_TARGET,
            "body_pages_hard_max": BODY_PAGES_HARD_MAX,
            "ai_section_pages_max": AI_SECTION_PAGES_MAX,
            "ai_section_chars_max": AI_TALENT_BODY_SUMMARY_MAX_CHARS,
        },
        "pdf_bytes": pdf_bytes,
        "min_section_chars": min_section,
        "min_pdf_bytes": min_pdf,
        "ok": not reasons,
        "reasons": reasons,
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", default=str(DEFAULT_DOCX))
    ap.add_argument("--pdf", default=str(DEFAULT_PDF))
    ap.add_argument("--min-section-chars", type=int, default=MIN_SECTION_CHARS)
    ap.add_argument("--min-pdf-bytes", type=int, default=MIN_PDF_BYTES)
    ap.add_argument("--attest-g1", action="store_true", help="Attest G1 only when verify ok")
    args = ap.parse_args()

    docx = Path(args.docx)
    pdf = Path(args.pdf)
    if docx.is_file() and pdf == DEFAULT_PDF:
        sibling = docx.with_suffix(".pdf")
        if sibling.is_file():
            pdf = sibling
    if META.is_file() and not docx.is_file():
        meta = json.loads(META.read_text(encoding="utf-8-sig"))
        outs = meta.get("outputs") or {}
        docx = ROOT / str(outs.get("docx", docx))
        pdf = ROOT / str(outs.get("pdf", pdf))

    report = verify(
        docx=docx,
        pdf=pdf,
        min_section=args.min_section_chars,
        min_pdf=args.min_pdf_bytes,
    )
    OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))

    if args.attest_g1 and report["ok"]:
        import subprocess
        import sys

        rc = subprocess.call(
            [
                sys.executable,
                str(ROOT / "scripts/attest_kstartup_startup_package_ai_human_gates_v1.py"),
                "--gate",
                "G1_form_mapping",
                "--status",
                "pass",
                "--note",
                f"filled_plan verify ok sections={len(report.get('section_chars', {}))} pdf_bytes={report.get('pdf_bytes')}",
            ],
            cwd=str(ROOT),
        )
        return rc

    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
