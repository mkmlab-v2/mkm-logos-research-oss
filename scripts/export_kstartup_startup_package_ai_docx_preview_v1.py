#!/usr/bin/env python3
"""Audit filled 도약 docx body flow for agent/human visual pre-attest (no GUI)."""
from __future__ import annotations

import argparse
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DOCX = ROOT / "reports/kstartup_startup_package_ai_filled/doyak_plan_filled_v1.docx"
OUT_META = ROOT / "reports/kstartup_startup_package_ai_filled_body_audit_latest.json"

MARKERS = (
    "문제 인식",
    "창업 아이템의 필요성",
    "개발 필요성",
    "개발 아이템",
    "실현 가능성",
    "성장전략",
    "아이템 개요",
)


def _utc() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _paragraph_xml_text(paragraph) -> str:
    return "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", paragraph._element.xml))


def audit_docx(docx: Path) -> dict[str, Any]:
    from docx import Document

    if not docx.is_file():
        return {"ok": False, "reason": "docx_missing"}

    doc = Document(str(docx))
    body_samples: list[dict[str, Any]] = []
    issues: list[str] = []

    for i, p in enumerate(doc.paragraphs):
        t = (p.text or "").strip()
        xt = _paragraph_xml_text(p).strip()
        blue = "0000FF" in p._element.xml.upper()
        if not t and not xt:
            continue
        if any(k in t or k in xt for k in MARKERS) or t.startswith("○") or t.startswith("ㅇ"):
            body_samples.append(
                {
                    "index": i,
                    "text": (t or xt)[:240],
                    "blue_xml": blue and ("※" in xt or "개발하고자" in xt),
                    "empty_o": bool(re.fullmatch(r"ㅇ\s*", t)),
                }
            )
        if blue and xt and ("※" in xt or "개발하고자" in xt):
            issues.append(f"P{i}:blue_instruction")
        if re.fullmatch(r"ㅇ\s*", t):
            issues.append(f"P{i}:empty_o")

    section_rows: dict[str, int] = {}
    for table in doc.tables:
        for row in table.rows:
            label = (row.cells[0].text or "").replace("\n", " ").strip()
            if len(row.cells) < 2:
                continue
            for key in ("아이템 개요", "문제 인식", "실현 가능성", "성장전략", "팀 구성", "AI 인재 활용"):
                if key in label:
                    section_rows[key] = len((row.cells[1].text or "").strip())

    return {
        "ok": not issues,
        "issues": issues,
        "body_samples": body_samples[:40],
        "section_table_chars": section_rows,
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", default=str(DEFAULT_DOCX))
    args = ap.parse_args()

    docx = Path(args.docx)
    if not docx.is_absolute():
        docx = (ROOT / docx).resolve()
    report = {
        "schema": "kstartup_startup_package_ai_filled_body_audit_v1",
        "generated_at_utc": _utc(),
        "docx": str(docx.relative_to(ROOT)).replace("\\", "/") if docx.is_file() else str(docx),
        "audit": audit_docx(docx),
    }
    OUT_META.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["audit"].get("ok") else 1


if __name__ == "__main__":
    raise SystemExit(main())
