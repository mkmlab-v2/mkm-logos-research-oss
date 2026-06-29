#!/usr/bin/env python3
"""Build NVIDIA Inception pitch deck (PDF + PPTX) for mkmlab / jema-ai.com."""
from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
OUT_PPTX = ROOT / "reports" / "nvidia_inception_pitch_deck_v1.pptx"
OUT_PDF = ROOT / "reports" / "nvidia_inception_pitch_deck_v1.pdf"
OUT_META = ROOT / "reports" / "nvidia_inception_pitch_deck_v1_latest.json"

SLIDES: list[tuple[str, list[str]]] = [
    (
        "MKM Lab (mkmlab)",
        [
            "Enterprise Document Intelligence for Regulated B2B",
            "https://jema-ai.com",
            "NVIDIA Inception Application — 2026",
        ],
    ),
    (
        "Problem",
        [
            "Teams store large document corpora but struggle to retrieve answers quickly.",
            "Generic chat tools lack audit-friendly, schema-first pipelines.",
            "Storage and inference costs grow without reproducible benchmarks.",
        ],
    ),
    (
        "Solution",
        [
            "Deterministic multi-lens document compression + RAG Q&A.",
            "Hybrid / on-prem friendly architecture for regulated workflows.",
            "Research prototypes isolated from production automation.",
        ],
    ),
    (
        "Product",
        [
            "MKM Document Intelligence Platform (Track C R&D).",
            "Ingest → compress/index → retrieve → explain with citations.",
            "Public site: jema-ai.com; benchmarks run locally with verified artifacts.",
        ],
    ),
    (
        "Target Market",
        [
            "Professional services and B2B operators with compliance needs.",
            "Healthcare-adjacent workflows (wellness/education copy only; not clinical claims).",
            "Technology partners needing OEM-style document intelligence.",
        ],
    ),
    (
        "Differentiators",
        [
            "Reproducible local benchmarks and schema-first pipelines.",
            "Clear separation: research [HYPO] vs production-ready modules.",
            "Public-facing claims tied to verified artifacts only.",
        ],
    ),
    (
        "NVIDIA Technology",
        [
            "Today: RTX-class GPUs for inference prototyping and RAG evaluation.",
            "Next: Nemotron and multimodal RAG benchmarks at scale via Inception.",
            "Interest: Innovation Lab for model optimization and retrieval quality.",
        ],
    ),
    (
        "Company",
        [
            "Incorporated 2021 · Korea (Republic of) · Gyeonggi.",
            "Brand: mkmlab · Product hub: jema-ai.com.",
            "Self-funded; small engineering team.",
        ],
    ),
    (
        "Traction",
        [
            "Live product website and document intelligence prototypes.",
            "Local GPU smoke tests and RAG evaluation harness in progress.",
            "B2B outreach and partner conversations (pre-revenue R&D stage).",
        ],
    ),
    (
        "Roadmap & Ask",
        [
            "Join NVIDIA Inception to validate RAG/Nemotron workloads at scale.",
            "Request: partner credits, Innovation Lab access, GTM guidance.",
            "Contact: moksorinw@no1kmedi.com",
        ],
    ),
]


def _build_pptx(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for title, bullets in SLIDES:
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(12.0), Inches(1.0))
        tf = box.text_frame
        tf.text = title
        p = tf.paragraphs[0]
        p.font.size = Pt(32)
        p.font.bold = True
        body = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.5), Inches(5.5))
        bf = body.text_frame
        bf.word_wrap = True
        for i, line in enumerate(bullets):
            para = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
            para.text = f"• {line}"
            para.font.size = Pt(20)
            para.space_after = Pt(10)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))


def _wrap(text: str, width: int = 95) -> list[str]:
    words = text.split()
    lines: list[str] = []
    cur: list[str] = []
    for w in words:
        trial = (" ".join(cur + [w])).strip()
        if len(trial) <= width:
            cur.append(w)
        else:
            if cur:
                lines.append(" ".join(cur))
            cur = [w]
    if cur:
        lines.append(" ".join(cur))
    return lines or [text]


def _build_pdf(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    w, h = landscape(letter)
    c = canvas.Canvas(str(path), pagesize=landscape(letter))
    for idx, (title, bullets) in enumerate(SLIDES, start=1):
        c.setFont("Helvetica-Bold", 28)
        c.drawString(0.75 * inch, h - 1.0 * inch, title)
        c.setFont("Helvetica", 14)
        y = h - 1.6 * inch
        for bullet in bullets:
            for line in _wrap(bullet):
                c.drawString(0.95 * inch, y, f"• {line}" if line == _wrap(bullet)[0] else f"  {line}")
                y -= 0.28 * inch
            y -= 0.08 * inch
        c.setFont("Helvetica", 10)
        c.drawRightString(w - 0.5 * inch, 0.45 * inch, f"{idx} / {len(SLIDES)}")
        c.showPage()
    c.save()


def main() -> int:
    _build_pptx(OUT_PPTX)
    _build_pdf(OUT_PDF)
    meta = {
        "schema": "nvidia_inception_pitch_deck_v1",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "pptx": str(OUT_PPTX),
        "pdf": str(OUT_PDF),
        "slide_count": len(SLIDES),
        "upload_hint": "NVIDIA Inception: Upload Company Pitch Deck — prefer PDF",
        "disclaimer": "Review before submit; no live-trading or unverified KPI claims.",
    }
    OUT_META.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(meta, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
